from __future__ import annotations

import fitz
import pytest
from docx import Document as DocxDocument
from openpyxl import Workbook

from app.ingest.parsers import (
    _ensure_public_url,
    parse_csv,
    parse_docx,
    parse_excel,
    parse_html,
    parse_image,
    parse_pdf,
    parse_text,
)
from app.llm.dashscope import clean_ocr_text


def test_markdown_heading_paths(tmp_path):
    path = tmp_path / "policy.md"
    path.write_text(
        "# 售后\n## 保修\n保修 1 年。\n## 退货\n7 天无理由退货。",
        encoding="utf-8",
    )
    parsed = parse_text(path)
    assert [block.heading_path for block in parsed.blocks] == ["售后 › 保修", "售后 › 退货"]
    assert "1 年" in parsed.blocks[0].text
    assert "7 天" in parsed.blocks[1].text


def test_csv_parser_keeps_sheet_and_row_locators(tmp_path):
    path = tmp_path / "prices.csv"
    path.write_text("名称,价格,库存\n星辰 X9,8888,12\n星光 P11,2199,30\n", encoding="utf-8-sig")
    parsed = parse_csv(path)
    assert parsed.blocks
    block = parsed.blocks[0]
    assert block.sheet_name == "CSV"
    assert block.row_start == 2 and block.row_end == 3
    assert "价格: 8888" in block.text
    assert "库存: 30" in block.text


def test_excel_parser_reads_sheet_names(tmp_path):
    path = tmp_path / "products.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "价格表"
    ws.append(["型号", "价格", "保修"])
    ws.append(["星辰 X9", 8888, "3 年"])
    wb.save(path)
    parsed = parse_excel(path)
    assert parsed.blocks
    assert parsed.blocks[0].sheet_name == "价格表"
    assert "星辰 X9" in parsed.blocks[0].text
    assert parsed.blocks[0].row_start == 2


def test_docx_parser_reads_headings_and_tables(tmp_path):
    path = tmp_path / "guide.docx"
    doc = DocxDocument()
    doc.add_heading("售后服务", level=1)
    doc.add_paragraph("主机保修 3 年。")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "项目"
    table.cell(0, 1).text = "内容"
    table.cell(1, 0).text = "保修"
    table.cell(1, 1).text = "3 年"
    doc.save(path)

    parsed = parse_docx(path)
    texts = "\n".join(block.text for block in parsed.blocks)
    assert "主机保修 3 年" in texts
    assert "保修 | 3 年" in texts
    assert any(block.heading_path == "售后服务" for block in parsed.blocks)


def test_pdf_parser_extracts_chinese_text(tmp_path):
    path = tmp_path / "policy.pdf"
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_textbox(
        fitz.Rect(40, 40, 560, 780),
        "星辰 X9 保修政策\n整机保修 3 年。\n七天无理由退货。",
        fontname="china-s",
        fontsize=12,
        lineheight=1.6,
    )
    pdf.save(path)
    pdf.close()

    parsed = parse_pdf(path)
    assert parsed.page_count == 1
    assert parsed.blocks and parsed.blocks[0].page_no == 1
    assert "保修" in parsed.blocks[0].text


def test_html_parser_extracts_article():
    html = (
        "<html><head><title>测试商品页</title></head><body><article>"
        "<h1>星光平板 P11</h1><p>售价 2199 元，整机保修 1 年。</p>"
        "</article></body></html>"
    )
    parsed = parse_html(html, fallback_title="测试商品页")
    text = "\n".join(block.text for block in parsed.blocks)
    assert "2199" in text and "保修" in text
    # trafilatura may prefer the article <h1> over <title>; both are valid titles.
    assert parsed.title in {"测试商品页", "星光平板 P11"}


def test_image_parser_requires_api_key(tmp_path):
    path = tmp_path / "image.png"
    path.write_bytes(b"\x89PNG\r\n\x1a\nfake")
    with pytest.raises(ValueError, match="DASHSCOPE_API_KEY"):
        parse_image(path)


def test_ocr_html_cleanup():
    raw = "```html\n<html><body><p>星光平板 P11</p><p>售价 2199 元</p></body></html>\n```"
    cleaned = clean_ocr_text(raw)
    assert "<" not in cleaned and ">" not in cleaned
    assert "星光平板 P11" in cleaned
    assert "售价 2199 元" in cleaned


@pytest.mark.parametrize(
    "url",
    [
        "http://localhost:8000/",
        "http://127.0.0.1/",
        "http://10.0.0.1/",
        "http://[::1]/",
        "file:///etc/passwd",
    ],
)
def test_ssrf_guard_rejects_private_or_invalid_urls(url: str):
    with pytest.raises(ValueError):
        _ensure_public_url(url)
