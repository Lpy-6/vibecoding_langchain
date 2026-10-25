"""
Document parsers: PDF (with scanned-page OCR), DOCX, TXT/MD, Excel/CSV, URL, image.

Each parser returns ParsedDocument with locator metadata so citations can point
to a page, worksheet/row range or heading path.
"""
from __future__ import annotations

import asyncio
import csv
import io
import ipaddress
import logging
import socket
from dataclasses import dataclass, field
from pathlib import Path
from urllib.parse import urlparse

import fitz
import httpx
from bs4 import BeautifulSoup
from docx import Document as DocxDocument
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook

from app.core.config import settings
from app.llm import dashscope

logger = logging.getLogger(__name__)


@dataclass
class ParsedBlock:
    text: str
    page_no: int | None = None
    sheet_name: str | None = None
    row_start: int | None = None
    row_end: int | None = None
    heading_path: str | None = None


@dataclass
class ParsedDocument:
    blocks: list[ParsedBlock] = field(default_factory=list)
    page_count: int = 0
    ocr_pages: int = 0
    title: str = ""
    meta: dict = field(default_factory=dict)


def _run_async(coro):
    try:
        asyncio.get_running_loop()
    except RuntimeError:
        return asyncio.run(coro)
    # Inside an existing loop (should not happen in worker threads): run in a new thread.
    import concurrent.futures

    with concurrent.futures.ThreadPoolExecutor(max_workers=1) as pool:
        return pool.submit(lambda: asyncio.run(coro)).result()


def parse_pdf(path: Path) -> ParsedDocument:
    result = ParsedDocument()
    doc = fitz.open(path)
    result.page_count = doc.page_count
    for page_index, page in enumerate(doc):
        text = page.get_text("text").strip()
        needs_ocr = len(text) < 30 and bool(page.get_images(full=True))
        if needs_ocr:
            if not settings.use_dashscope:
                continue
            try:
                pix = page.get_pixmap(dpi=200)
                text = _run_async(dashscope.ocr_image(pix.tobytes("png"), "image/png"))
                result.ocr_pages += 1
            except Exception as exc:  # noqa: BLE001
                logger.warning("OCR failed page=%s: %s", page_index + 1, exc)
                text = ""
        if text.strip():
            result.blocks.append(ParsedBlock(text=text, page_no=page_index + 1))
    doc.close()
    result.meta["parser"] = "pymupdf"
    if not result.blocks:
        raise ValueError("PDF 未提取到文本；若为扫描件，请在 .env 配置 DASHSCOPE_API_KEY 以启用 OCR")
    return result


def parse_docx(path: Path) -> ParsedDocument:
    result = ParsedDocument()
    doc = DocxDocument(str(path))
    headings: list[str] = []
    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = Paragraph(child, doc)
            text = para.text.strip()
            if not text:
                continue
            style = (para.style.name or "").lower()
            if "heading" in style or style.startswith("标题"):
                level = 1
                for ch in style:
                    if ch.isdigit():
                        level = int(ch)
                        break
                headings = headings[: max(0, level - 1)] + [text]
                continue
            result.blocks.append(ParsedBlock(text=text, heading_path=" › ".join(headings) if headings else None))
        elif tag == "tbl":
            table = Table(child, doc)
            lines: list[str] = []
            for row in table.rows:
                cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                if any(cells):
                    lines.append(" | ".join(f"{c}" for c in cells))
            if lines:
                result.blocks.append(
                    ParsedBlock(text="\n".join(lines), heading_path=" › ".join(headings) if headings else None)
                )
    result.meta["parser"] = "python-docx"
    result.page_count = 0
    if not result.blocks:
        raise ValueError("Word 文档中没有可提取的文本")
    return result


def parse_text(path: Path) -> ParsedDocument:
    raw = path.read_bytes()
    text = ""
    for enc in ("utf-8-sig", "utf-8", "gb18030", "utf-16"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if not text:
        raise ValueError("无法识别文本编码")
    result = ParsedDocument()
    headings: list[str] = []
    buffer: list[str] = []

    def flush() -> None:
        content = "\n".join(buffer).strip()
        buffer.clear()
        if content:
            result.blocks.append(ParsedBlock(text=content, heading_path=" › ".join(headings) if headings else None))

    for line in text.splitlines():
        stripped = line.strip()
        if stripped.startswith("#") and path.suffix.lower() in {".md", ".markdown"}:
            flush()
            level = len(stripped) - len(stripped.lstrip("#"))
            headings = headings[: max(0, level - 1)] + [stripped.lstrip("#").strip()]
            continue
        buffer.append(line)
    flush()
    result.meta["parser"] = "text"
    if not result.blocks:
        raise ValueError("文件中没有可提取的文本")
    return result


def _sheet_blocks(rows: list[list[str]], sheet_name: str, rows_per_chunk: int = 20) -> list[ParsedBlock]:
    if not rows:
        return []
    header = rows[0]
    data_rows = rows[1:]
    blocks: list[ParsedBlock] = []
    for start in range(0, len(data_rows), rows_per_chunk):
        chunk = data_rows[start : start + rows_per_chunk]
        lines = [" | ".join(str(h).strip() for h in header)]
        for row in chunk:
            pairs = []
            for idx, value in enumerate(row):
                value = str(value).strip()
                if value == "" or value == "None":
                    continue
                key = header[idx] if idx < len(header) and str(header[idx]).strip() else f"列{idx + 1}"
                pairs.append(f"{key}: {value}")
            if pairs:
                lines.append("；".join(pairs))
        if len(lines) > 1:
            blocks.append(
                ParsedBlock(
                    text="\n".join(lines),
                    sheet_name=sheet_name,
                    row_start=start + 2,
                    row_end=start + 1 + len(chunk),
                )
            )
    return blocks


def parse_excel(path: Path) -> ParsedDocument:
    result = ParsedDocument()
    wb = load_workbook(path, read_only=True, data_only=True)
    for ws in wb.worksheets:
        rows: list[list[str]] = []
        for row in ws.iter_rows(values_only=True):
            values = ["" if cell is None else str(cell) for cell in row]
            if any(v.strip() for v in values):
                rows.append(values)
        result.blocks.extend(_sheet_blocks(rows, ws.title))
    wb.close()
    result.meta["parser"] = "openpyxl"
    if not result.blocks:
        raise ValueError("表格中没有可提取的数据")
    return result


def parse_csv(path: Path) -> ParsedDocument:
    raw = path.read_bytes()
    text = None
    for enc in ("utf-8-sig", "utf-8", "gb18030", "utf-16"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise ValueError("无法识别 CSV 编码")
    rows = [row for row in csv.reader(io.StringIO(text)) if any(str(c).strip() for c in row)]
    result = ParsedDocument()
    result.blocks = _sheet_blocks(rows, "CSV")
    result.meta["parser"] = "csv"
    if not result.blocks:
        raise ValueError("CSV 中没有可提取的数据")
    return result


def parse_image(path: Path, mime_type: str = "image/png") -> ParsedDocument:
    if not settings.use_dashscope:
        raise ValueError("图片 OCR 需要配置 DASHSCOPE_API_KEY")
    data = path.read_bytes()
    text = _run_async(dashscope.ocr_image(data, mime_type))
    result = ParsedDocument(page_count=1, ocr_pages=1, meta={"parser": "qwen-vl-ocr"})
    if text.strip():
        result.blocks.append(ParsedBlock(text=text.strip(), page_no=1))
    else:
        raise ValueError("OCR 未识别到文字")
    return result


# ---------------- URL ----------------
def _ensure_public_url(url: str) -> None:
    parsed = urlparse(url)
    if parsed.scheme not in {"http", "https"}:
        raise ValueError("仅支持 http/https 链接")
    host = parsed.hostname
    if not host:
        raise ValueError("链接缺少主机名")
    if host.lower() in {"localhost", "localhost.localdomain"}:
        raise ValueError("不允许访问本机地址")
    try:
        infos = socket.getaddrinfo(host, None)
    except socket.gaierror as exc:
        raise ValueError(f"域名解析失败: {host}") from exc
    for info in infos:
        ip = ipaddress.ip_address(info[4][0])
        if ip.is_private or ip.is_loopback or ip.is_link_local or ip.is_reserved or ip.is_multicast:
            raise ValueError("不允许访问内网或保留地址")


def fetch_url(url: str, max_bytes: int = 5 * 1024 * 1024) -> tuple[str, str]:
    current = url
    headers = {"User-Agent": "Mozilla/5.0 (compatible; RagKnowledgeBot/1.0)"}
    with httpx.Client(timeout=20, follow_redirects=False, headers=headers) as client:
        for _ in range(4):
            _ensure_public_url(current)
            with client.stream("GET", current) as resp:
                if resp.status_code in {301, 302, 303, 307, 308}:
                    location = resp.headers.get("location")
                    if not location:
                        raise ValueError("重定向缺少目标地址")
                    current = str(resp.url.join(location))
                    continue
                resp.raise_for_status()
                chunks: list[bytes] = []
                size = 0
                for chunk in resp.iter_bytes():
                    size += len(chunk)
                    if size > max_bytes:
                        raise ValueError("网页内容超过 5MB 限制")
                    chunks.append(chunk)
                encoding = resp.encoding or "utf-8"
                return b"".join(chunks).decode(encoding, errors="replace"), str(resp.url)
    raise ValueError("重定向次数过多")


def parse_html(html: str, fallback_title: str = "") -> ParsedDocument:
    title = fallback_title
    content = ""
    try:
        import trafilatura
        from trafilatura import extract, extract_metadata

        meta = extract_metadata(html)
        if meta and meta.title:
            title = meta.title
        content = extract(
            html,
            include_tables=True,
            include_comments=False,
            favor_recall=True,
            output_format="markdown",
        ) or ""
    except Exception as exc:  # noqa: BLE001
        logger.warning("trafilatura extract failed: %s", exc)
    if not content.strip():
        soup = BeautifulSoup(html, "html.parser")
        if not title:
            title = soup.title.get_text(strip=True) if soup.title else ""
        for tag in soup(["script", "style", "nav", "footer", "header"]):
            tag.decompose()
        content = soup.get_text("\n", strip=True)
    if not content.strip():
        raise ValueError("网页正文为空或无法解析")
    result = ParsedDocument(title=title.strip(), page_count=0, meta={"parser": "trafilatura"})
    paragraphs = [p.strip() for p in content.split("\n\n") if p.strip()]
    for para in paragraphs:
        result.blocks.append(ParsedBlock(text=para))
    return result


def parse_path(path: Path, mime_type: str = "") -> ParsedDocument:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return parse_pdf(path)
    if suffix in {".docx"}:
        return parse_docx(path)
    if suffix in {".txt", ".md", ".markdown"}:
        return parse_text(path)
    if suffix in {".xlsx", ".xlsm"}:
        return parse_excel(path)
    if suffix == ".csv":
        return parse_csv(path)
    if suffix in {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tiff"}:
        return parse_image(path, mime_type or "image/png")
    if suffix == ".doc":
        raise ValueError("暂不支持旧版 .doc，请转换为 .docx 后上传")
    if suffix == ".xls":
        raise ValueError("暂不支持旧版 .xls，请转换为 .xlsx 后上传")
    raise ValueError(f"不支持的文件类型: {suffix or 'unknown'}")
